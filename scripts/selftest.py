#!/usr/bin/env python3
"""Regression checks for word_gov_format_cn.py.

Usage: python scripts/selftest.py
Exits 0 when all checks pass.
"""
import base64
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Cm, Inches

SCRIPT = Path(__file__).with_name("word_gov_format_cn.py")
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFAAH/q842iQAAAABJRU5ErkJggg=="
)

failures = []

EAST_ASIA = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"


def east_asia_font(paragraph):
    for run in paragraph.runs:
        rpr = run._element.rPr
        if rpr is not None and rpr.rFonts is not None:
            return rpr.rFonts.get(EAST_ASIA)
    return None


def check(name, condition, detail=""):
    if condition:
        print(f"  ok   {name}")
    else:
        print(f"  FAIL {name} {detail}")
        failures.append(name)


def run(*args):
    return subprocess.run([sys.executable, str(SCRIPT), *args], capture_output=True, text=True)


def build_doc(path: Path):
    png = path.parent / "px.png"
    png.write_bytes(PNG)
    d = Document()
    d.add_paragraph("某某市某某局")
    d.add_paragraph("关于开展公文格式自查工作的通知")
    d.add_paragraph("各科室,各单位:")
    d.add_paragraph("访问 www.hangzhou.gov.cn ,邮箱 test@abc.com 可咨询.")
    d.add_paragraph("Excel 2019 and e.g. Word, version 1.2.3. 时间 10:30.")
    d.add_paragraph("一、自查范围")
    d.add_paragraph("1.版式要素是否齐全")
    p = d.add_paragraph("正文前段,")
    p.add_run("重点强调文字,").bold = True
    p.add_run("后段.")
    d.add_paragraph("保留图片:")
    d.add_paragraph("他说\"没问题\"就走了. don't panic.")
    d.add_paragraph("").add_run().add_picture(str(png), width=Inches(0.4))
    d.save(str(path))


def main():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        src = root / "in" / "样例.docx"
        src.parent.mkdir()
        build_doc(src)

        print("layout:")
        out = root / "out"
        r = run(str(src), "-o", str(out))
        check("run exits 0", r.returncode == 0, r.stderr)
        doc = Document(str(out / "样例_排版.docx"))
        s = doc.sections[0]
        check("A4 page", abs(s.page_width.cm - 21) < 0.05 and abs(s.page_height.cm - 29.7) < 0.05)
        check(
            "margins 3.7/3.5/2.8/2.6",
            [round(m.cm, 2) for m in (s.top_margin, s.bottom_margin, s.left_margin, s.right_margin)]
            == [3.7, 3.5, 2.8, 2.6],
        )
        fonts = {p.text.strip()[:12]: east_asia_font(p) for p in doc.paragraphs if p.runs}
        check("issuer font", fonts.get("某某市某某局") == "方正小标宋简体")
        check("level1 font", fonts.get("一、自查范围") == "黑体")
        check("body spacing", round(doc.paragraphs[2].paragraph_format.line_spacing.pt, 1) == 28.0)
        check("body indent 32pt", round(doc.paragraphs[2].paragraph_format.first_line_indent.pt, 1) == 32.0)

        print("punctuation:")
        texts = [p.text for p in doc.paragraphs]
        punct = next(t for t in texts if "hangzhou" in t)
        check("url protected", "www.hangzhou.gov.cn" in punct, punct)
        check("email protected", "test@abc.com" in punct, punct)
        check("cjk comma converted", "，" in punct, punct)
        latin = next(t for t in texts if "Excel" in t)
        check("e.g. protected", "e.g." in latin, latin)
        check("version protected", "1.2.3" in latin, latin)
        check("time protected", "10:30" in latin, latin)
        check("list numbering kept", any(t.startswith("1.") for t in texts), texts)
        check("chinese quotes", any("“" in t for t in texts), texts)
        check("ascii apostrophe kept", any("don't" in t for t in texts), texts)

        print("content fidelity:")
        check("picture preserved", len(doc.inline_shapes) == 1, str(len(doc.inline_shapes)))
        bold = [p for p in doc.paragraphs if "重点强调" in p.text]
        check("inline bold preserved", bool(bold) and any(r.bold for r in bold[0].runs))

        print("safety:")
        r = run(str(src), "-o", str(out))
        check("refuses overwrite", r.returncode == 2, r.stdout + r.stderr)
        r = run(str(src), "-o", str(out), "--force")
        check("--force overwrites", r.returncode == 0, r.stderr)
        r = run(str(src), "-o", str(root / "dry"), "--dry-run")
        check("dry-run writes nothing", not (root / "dry").exists() and r.returncode == 0)

        copy = root / "in" / "原件.docx"
        copy.write_bytes(src.read_bytes())
        before = copy.read_bytes()
        r = run(str(copy), "--in-place")
        backups = list(copy.parent.glob("原件.docx.bak-*"))
        check("in-place creates backup", len(backups) == 1, str(backups))
        check("backup is original", backups and backups[0].read_bytes() == before)

        many = root / "many"
        many.mkdir()
        (many / "a.docx").write_bytes(src.read_bytes())
        (many / "b.docx").write_bytes(src.read_bytes())
        r = run(str(many), "-o", str(root / "one.docx"))
        check("batch to single file rejected", r.returncode == 1, r.stderr)

        (many / "broken.docx").write_text("not a zip", encoding="utf-8")
        r = run(str(many), "-o", str(root / "batchout"))
        produced = sorted(p.name for p in (root / "batchout").glob("*.docx"))
        check("corrupt input skipped, batch continues", produced == ["a_排版.docx", "b_排版.docx"], str(produced))
        check("corrupt input sets exit 1", r.returncode == 1)

    print()
    if failures:
        print(f"{len(failures)} check(s) failed: {', '.join(failures)}")
        return 1
    print("all checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
