#!/usr/bin/env python3
import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except Exception:
    print("缺少依赖：python-docx，请先安装：pip install python-docx", file=sys.stderr)
    raise


LEVEL1_RE = re.compile(r"^\s*[一二三四五六七八九十百]+、")
LEVEL2_RE = re.compile(r"^\s*（[一二三四五六七八九十百]+）")
LEVEL3_RE = re.compile(r"^\s*[0-9]+[.]")
LEVEL4_RE = re.compile(r"^\s*（[0-9]+）")

FULLWIDTH_MAP = {
    ",": "，",
    ";": "；",
    ":": "：",
    "!": "！",
    "?": "？",
    "(": "（",
    ")": "）",
    "[": "【",
    "]": "】",
}


def normalize_punctuation(text: str) -> str:
    chars = list(text)
    out = []
    dquote_open = True   # double quote pairing state
    squote_open = True   # single quote pairing state
    for i, ch in enumerate(chars):
        if ch == ".":
            prev_digit = i > 0 and chars[i - 1].isdigit()
            next_digit = i + 1 < len(chars) and chars[i + 1].isdigit()
            out.append("." if (prev_digit and next_digit) else "。")
            continue
        if ch == ",":
            prev_digit = i > 0 and chars[i - 1].isdigit()
            next_digit = i + 1 < len(chars) and chars[i + 1].isdigit()
            out.append("," if (prev_digit and next_digit) else "，")
            continue
        # Convert straight double quotes to Chinese curly quotes (pairing)
        if ch == '"':
            out.append("\u201c" if dquote_open else "\u201d")
            dquote_open = not dquote_open
            continue
        # Convert straight single quotes to Chinese single quotes
        if ch == "'":
            out.append("\u2018" if squote_open else "\u2019")
            squote_open = not squote_open
            continue
        out.append(FULLWIDTH_MAP.get(ch, ch))
    return "".join(out)


def set_run_font(run, cn_font: str, en_font: str, size_pt: float, bold: bool = False, color=None):
    run.font.name = en_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn_font)
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_spacing(p, line_pt: float, first_line_indent_pt: float = 0):
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(line_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)


def classify_paragraph(text: str) -> str:
    if LEVEL1_RE.match(text):
        return "level1"
    if LEVEL2_RE.match(text):
        return "level2"
    if LEVEL3_RE.match(text):
        return "level3"
    if LEVEL4_RE.match(text):
        return "level4"
    return "body"


def format_paragraph(
    p,
    en_font: str,
    is_issuer: bool = False,
    is_title: bool = False,
):
    text = p.text.strip()
    if not text:
        return

    p.text = normalize_punctuation(p.text)
    text = p.text.strip()

    # 发文机关标志：方正小标宋简体，二号，红色
    if is_issuer:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_line = 32 if len(text) > 28 else 38
        set_paragraph_spacing(p, line_pt=title_line, first_line_indent_pt=0)
        for r in p.runs:
            set_run_font(
                r,
                cn_font="方正小标宋简体",
                en_font=en_font,
                size_pt=22,
                bold=False,
                color=RGBColor(0xFF, 0x00, 0x00),
            )
        return

    # 公文标题：方正小标宋简体，二号
    if is_title:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_line = 32 if len(text) > 28 else 38
        set_paragraph_spacing(p, line_pt=title_line, first_line_indent_pt=0)
        for r in p.runs:
            set_run_font(r, "方正小标宋简体", en_font, 22, bold=False)
        return

    style = classify_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if style == "level1":
        set_paragraph_spacing(p, line_pt=28, first_line_indent_pt=0)
        for r in p.runs:
            set_run_font(r, "黑体", en_font, 16, bold=False)
    elif style == "level2":
        set_paragraph_spacing(p, line_pt=28, first_line_indent_pt=0)
        for r in p.runs:
            set_run_font(r, "楷体_GB2312", en_font, 16, bold=False)
    elif style in ("level3", "level4"):
        set_paragraph_spacing(p, line_pt=28, first_line_indent_pt=0)
        for r in p.runs:
            set_run_font(r, "仿宋_GB2312", en_font, 16, bold=False)
    else:
        # 正文：仿宋_GB2312，三号，首行缩进2字
        set_paragraph_spacing(p, line_pt=28, first_line_indent_pt=32)
        for r in p.runs:
            set_run_font(r, "仿宋_GB2312", en_font, 16, bold=False)


def format_document(
    file_path: Path,
    out_path: Path,
    en_font: str,
):
    doc = Document(str(file_path))

    # 页面设置（A4 + 指定边距，对应版心 156mm x 225mm）
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(3.7)
        section.bottom_margin = Cm(3.5)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.6)

    # 段落角色判定：
    # - 发文机关标志：以"局""部""委""办""中心"等结尾的独立短行
    # - 公文标题：第一个非空段落（若不像发文机关）或第二个非空段落
    non_empty = [(idx, p) for idx, p in enumerate(doc.paragraphs) if p.text.strip()]
    
    issuer_index = None
    title_index = None
    
    if non_empty:
        first_idx, first_p = non_empty[0]
        first_text = first_p.text.strip()
        # Check if first paragraph looks like an issuer mark
        issuer_suffixes = ('局', '部', '委', '办', '中心', '处', '室', '会', '厅')
        if len(first_text) <= 30 and first_text.endswith(issuer_suffixes) and '\n' not in first_text:
            issuer_index = first_idx
            if len(non_empty) >= 2:
                title_index = non_empty[1][0]
        else:
            title_index = first_idx
    
    # If no issuer found but there's a second short paragraph that looks like one
    if issuer_index is None and len(non_empty) >= 2:
        second_idx, second_p = non_empty[1]
        second_text = second_p.text.strip()
        issuer_suffixes = ('局', '部', '委', '办', '中心', '处', '室', '会', '厅')
        if len(second_text) <= 30 and second_text.endswith(issuer_suffixes) and '\n' not in second_text:
            issuer_index = second_idx
            title_index = non_empty[0][0]

    for idx, p in enumerate(doc.paragraphs):
        format_paragraph(
            p,
            en_font=en_font,
            is_issuer=(issuer_index is not None and idx == issuer_index),
            is_title=(title_index is not None and idx == title_index),
        )

    # 表格内容按正文处理
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    format_paragraph(
                        p,
                        en_font=en_font,
                        is_issuer=False,
                        is_title=False,
                    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def iter_docx(path: Path):
    if path.is_file() and path.suffix.lower() == ".docx":
        yield path
        return
    if path.is_dir():
        for p in sorted(path.glob("*.docx")):
            if p.name.startswith("~$"):
                continue
            yield p


def build_output_path(src: Path, output: Path | None, in_place: bool):
    if in_place:
        return src
    if output is None:
        return src.with_name(f"{src.stem}_排版.docx")
    if output.suffix.lower() == ".docx":
        return output
    return output / f"{src.stem}_排版.docx"


def main():
    parser = argparse.ArgumentParser(description="中文公文版式排版（.docx）")
    parser.add_argument("input", help="输入文件或目录（仅 .docx）")
    parser.add_argument("-o", "--output", help="输出文件或输出目录（可选）")
    parser.add_argument("--in-place", action="store_true", help="原地覆盖（谨慎）")
    parser.add_argument("--en-font", default="Times New Roman", help="阿拉伯数字/英文默认字体")
    args = parser.parse_args()

    in_path = Path(args.input).expanduser().resolve()
    out_arg = Path(args.output).expanduser().resolve() if args.output else None

    targets = list(iter_docx(in_path))
    if not targets:
        print("未找到可处理的 .docx 文件", file=sys.stderr)
        return 1

    for src in targets:
        out_path = build_output_path(src, out_arg, args.in_place)
        format_document(
            file_path=src,
            out_path=out_path,
            en_font=args.en_font,
        )
        print(f"已生成: {out_path}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
