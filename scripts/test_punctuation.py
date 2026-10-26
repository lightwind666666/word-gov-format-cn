#!/usr/bin/env python3
"""In-memory regression tests; run with python -B scripts/test_punctuation.py."""
import base64
import unittest
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches

from word_gov_format_cn import format_paragraph, rewrite_paragraph_text


class PunctuationTests(unittest.TestCase):
    def test_protected_text_across_run_boundaries(self):
        cases = (
            "\u7f51\u5740https://example.com/a?x=1&y=2",
            "\u90ae\u7bb1test@example.com",
            "\u65f6\u95f410:30:45",
            "\u57df\u540dwww.example.com",
            "\u793a\u4f8bdon't e.g. v1.2.3 file.docx",
            "  12.\u529e\u7406\u8981\u6c42",
        )
        document = Document()
        for text in cases:
            partitions = [
                [text[:split], text[split:]]
                for split in range(1, len(text))
            ] + [list(text)]
            for parts in partitions:
                with self.subTest(parts=parts):
                    paragraph = document.add_paragraph()
                    for part in parts:
                        paragraph.add_run(part).italic = True

                    format_paragraph(paragraph, "Times New Roman")

                    self.assertEqual(paragraph.text, text)
                    self.assertEqual(
                        [run.text for run in paragraph.runs], parts,
                    )
                    self.assertTrue(all(run.italic for run in paragraph.runs))

    def test_chinese_punctuation_and_quotes_across_runs(self):
        text = '\u4ed6\u8bf4"\u6ca1\u95ee\u9898",\u5404\u5355\u4f4d;\u529e\u7406.'
        expected = (
            "\u4ed6\u8bf4\u201c\u6ca1\u95ee\u9898\u201d\uff0c"
            "\u5404\u5355\u4f4d\uff1b\u529e\u7406\u3002"
        )
        document = Document()
        for split in range(1, len(text)):
            with self.subTest(split=split):
                paragraph = document.add_paragraph()
                paragraph.add_run(text[:split]).bold = True
                paragraph.add_run(text[split:]).italic = True
                format_paragraph(paragraph, "Times New Roman")
                self.assertEqual(paragraph.text, expected)
                self.assertEqual(paragraph.runs[0].text, expected[:split])
                self.assertEqual(paragraph.runs[1].text, expected[split:])
                self.assertTrue(paragraph.runs[0].bold)
                self.assertTrue(paragraph.runs[1].italic)

    def test_pure_english_is_unchanged(self):
        text = 'URL: https://example.com, "quoted"; don\'t (change) [this].'
        paragraph = Document().add_paragraph()
        for char in text:
            paragraph.add_run(char)
        format_paragraph(paragraph, "Times New Roman")
        self.assertEqual(paragraph.text, text)

    def test_disabled_normalization_is_unchanged(self):
        text = '\u4ed6\u8bf4"\u53ef\u4ee5",\u6536\u5230.'
        paragraph = Document().add_paragraph()
        for char in text:
            paragraph.add_run(char)
        format_paragraph(paragraph, "Times New Roman", normalize=False)
        self.assertEqual(paragraph.text, text)

    def test_hyperlink_preserved_and_following_runs_aligned(self):
        document = Document()
        paragraph = document.add_paragraph("\u53c2\u9605,\u94fe\u63a5 ")
        hyperlink = OxmlElement("w:hyperlink")
        relationship = document.part.relate_to(
            "https://example.com", RELATIONSHIP_TYPE.HYPERLINK, is_external=True,
        )
        hyperlink.set(qn("r:id"), relationship)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "example.com"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        paragraph.add_run(" \u65f6\u95f410")
        paragraph.add_run(":")
        paragraph.add_run("30 \u529e\u7406.")
        before = hyperlink.xml

        format_paragraph(paragraph, "Times New Roman")

        self.assertEqual(hyperlink.xml, before)
        self.assertEqual(
            paragraph.text,
            "\u53c2\u9605\uff0c\u94fe\u63a5 example.com "
            "\u65f6\u95f410:30 \u529e\u7406\u3002",
        )
        stream = BytesIO()
        document.save(stream)
        stream.seek(0)
        restored = Document(stream)
        self.assertEqual(restored.paragraphs[0].text, paragraph.text)
        self.assertEqual(
            restored.part.rels[relationship].target_ref, "https://example.com",
        )

    def test_complex_runs_preserved_and_offsets_kept(self):
        document = Document()
        paragraph = document.add_paragraph("\u524d\u6587,")
        complex_run = paragraph.add_run("\u539f\u6837,\t\u6362\u884c.\n")
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), "begin")
        complex_run._r.append(field)
        paragraph.add_run("\u540e\u6587")
        paragraph.add_run(",")
        paragraph.add_run("\u7ed3\u675f.")
        original = complex_run._r.xml

        rewrite_paragraph_text(paragraph, "body", True)

        self.assertEqual(complex_run._r.xml, original)
        self.assertEqual(
            paragraph.text,
            "\u524d\u6587\uff0c\u539f\u6837,\t\u6362\u884c.\n"
            "\u540e\u6587\uff0c\u7ed3\u675f\u3002",
        )

    def test_picture_and_bold_survive_round_trip(self):
        document = Document()
        paragraph = document.add_paragraph("\u524d\u6587,")
        picture_run = paragraph.add_run()
        png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8Dw"
            "HwAFAAH/q842iQAAAABJRU5ErkJggg=="
        )
        picture_run.add_picture(BytesIO(png), width=Inches(0.4))
        paragraph.add_run("\u540e\u6587.").bold = True
        picture_xml = picture_run._r.find(qn("w:drawing")).xml

        format_paragraph(paragraph, "Times New Roman")

        self.assertEqual(picture_run._r.find(qn("w:drawing")).xml, picture_xml)
        stream = BytesIO()
        document.save(stream)
        stream.seek(0)
        restored = Document(stream)
        self.assertEqual(len(restored.inline_shapes), 1)
        self.assertEqual(
            restored.paragraphs[0].text, "\u524d\u6587\uff0c\u540e\u6587\u3002",
        )
        self.assertTrue(restored.paragraphs[0].runs[-1].bold)

    def test_table_cell_uses_the_same_cross_run_protection(self):
        document = Document()
        paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        for part in ("\u90ae\u7bb1test@example", ".", "com \u8bf7\u8054\u7cfb."):
            paragraph.add_run(part)

        format_paragraph(paragraph, "Times New Roman")

        self.assertEqual(
            paragraph.text, "\u90ae\u7bb1test@example.com \u8bf7\u8054\u7cfb\u3002",
        )


if __name__ == "__main__":
    unittest.main()
